#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""店舗オープン日を集めて docs/open_dates.json（アプリ同梱の既定値）を生成する。

情報源（優先順）:
 1. lease: data/sources/グループ全社賃貸借契約_*.xlsx（5社全店・「店舗OPEN日初回営業開始日」「譲受店舗営業開始日」）
    店舗マスターとは電話番号で突合、無ければ郵便番号＋店名で突合。※機密ファイル。data/ は git 管理外。
 2. sales: data/sources/202702期売上.xlsx（ES 約160店・「オープン日(譲受日)」）… lease に無い店の補完
 3. news : https://www.eatstyle.jp/wp-json/wp/v2/news（ES の「○○店 オープン」記事。記事日付＝オープン日）
店舗マスター: https://hiiragi-hd.jp/wp-json/microcms-cache/v1/data/shop（アプリと同じ）

出力の JSON は日付と出典種別だけを含む（家賃条件等は含めない）。公開 repo の docs/ に置いてよい。
使い方: /usr/bin/python3 tools/fetch_open_dates.py   （openpyxl は CLT python3 のみ）
        新しい賃貸借一覧を入手したら data/sources/ に置いて再実行 → bump_version.py → commit/push。
アプリ側: 起動時に open_dates.json を既定値として読む。ユーザーがアプリで入力したオープン日が優先。
"""
import datetime
import glob
import html
import json
import re
import unicodedata
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "open_dates.json"
SRC = ROOT / "data" / "sources"
ES_NEWS = "https://www.eatstyle.jp/wp-json/wp/v2/news?per_page=100&page={}&_fields=id,date,title,link"
SHOP_API = "https://hiiragi-hd.jp/wp-json/microcms-cache/v1/data/shop?limit=100&offset={}"

BRAND_ALIAS = {"サンクゼール久世福商店": "久世福商店", "マヌカンピスケレス": "マヌカンピス・ケレス", "ブレッザテラス": "ブレッザカフェ"}
NAME_ALIAS = {"イオンモール鹿児島": "イオンモールKAGOSHIMABAY"}
# 売上一覧の略称 → 正規化名の置換（前方の略ブランド記号は除去してから適用）
SALES_ABBR = [("AM", "イオンモール"), ("AT", "イオンタウン"), ("RS", ""), ("PP", "パークプレイス"), ("IY", "イトーヨーカドー"),
              ("OL", "ジアウトレット"), ("らら", "ららぽーと"), ("ゆめ", "ゆめタウン"), ("アミュ", "アミュプラザ"), ("木の葉", "木の葉モール橋本")]


def get(url):
    with urllib.request.urlopen(url, timeout=30) as r:
        return json.load(r)


def nrm(s):
    s = unicodedata.normalize("NFKC", html.unescape(str(s or "")))
    s = s.replace("​", "")
    return re.sub(r"[\s・&店()（）]", "", s)


def core(name):
    x = re.sub(r"店$", "", nrm(name))
    return NAME_ALIAS.get(x, x)


def digits(s):
    return re.sub(r"\D", "", str(s or ""))


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    m = re.match(r"(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})", str(v or "").strip())
    return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else ""


def load_shops():
    shops, off = [], 0
    while True:
        d = get(SHOP_API.format(off))
        shops += d["contents"]
        off += 100
        if off >= d["totalCount"]:
            break
    return shops


def from_lease(shops):
    files = sorted(glob.glob(str(SRC / "グループ全社賃貸借契約_*.xlsx")))
    if not files:
        print("lease: data/sources/ に賃貸借契約一覧がありません（スキップ）")
        return {}, "なし"
    import openpyxl
    f = files[-1]
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    lease = []
    for ws in wb.worksheets:
        if ws.title == "見本":
            continue
        for r in ws.iter_rows(min_row=4, values_only=True):
            if not r[3] or not str(r[3]).strip():
                continue
            lease.append({"company": ws.title, "name": str(r[3]).strip(), "zip": str(r[5] or "").strip(), "phone": digits(r[7]),
                          "open": to_date(r[16]), "acq": to_date(r[19]), "close": to_date(r[22])})
    byphone = {}
    for l in lease:
        if l["phone"]:
            byphone.setdefault(l["phone"], []).append(l)
    out = {}
    for s in shops:
        cands = byphone.get(digits(s.get("phone")), [])
        m = None
        if len(cands) == 1:
            m = cands[0]
        elif len(cands) > 1:
            bn = [b["name"] for b in s.get("brands", [])]
            c2 = [c for c in cands if any(nrm(b)[:3] in nrm(c["name"]) for b in bn) or nrm(s["name"]) in nrm(c["name"])]
            m = (c2 or cands)[0]
        else:
            z = s.get("postal_code", "")
            c3 = [l for l in lease if l["zip"] == z and (nrm(s["name"]) in nrm(l["name"]) or nrm(l["name"]).endswith(nrm(s["name"])[-4:]))]
            if c3:
                m = c3[0]
        if m and m["open"]:
            note = f"譲受 {m['acq']}" if m["acq"] and m["acq"] != m["open"] else ("譲受店" if m["acq"] else "")
            out[s["id"]] = {"openDate": m["open"], "acquiredDate": m["acq"] or "", "note": note, "source": "lease"}
    return out, Path(f).name


def from_sales(shops):
    files = sorted(glob.glob(str(SRC / "*期売上*.xlsx")))
    if not files:
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(files[-1], read_only=True, data_only=True)
    sheets = [n for n in wb.sheetnames if "全店" in n]
    ws = wb[sheets[-1]]
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0] is None or not str(r[0]).strip().isdigit() or not r[2]:
            continue
        rows.append((str(r[1] or ""), str(r[2]), to_date(r[3])))
    out = {}
    for brand_s, name_s, d in rows:
        if not d:
            continue
        x = nrm(name_s)
        x = re.sub(r"^(31|ﾋﾞﾋﾞﾝ亭|ビビン亭|いきなり|ﾍﾟｯﾊﾟｰ|ペッパー|ﾋﾅﾀﾉ蔵|ヒナタノ蔵|BP|串ｶﾂ|串カツ|魁力屋|YOSHIMI|KF|SK|幸|ﾌﾞﾚｯｻﾞｶﾌｪ|ブレッザカフェ)", "", x)
        for a, b in SALES_ABBR:
            x = x.replace(a, b) if a in x else x
        cands = [s for s in shops if nrm(brand_s)[:2] in "".join(nrm(b["name"]) for b in s.get("brands", [])) or nrm(brand_s) in ("サンク久世福", "ＹＯＳＨＩＭＩ", "ヒナタノ蔵旧ジドリーノ", "ヒナタノ蔵旧勝手場ごった")]
        hit = [s for s in cands if x and (core(s["name"]) == x or x in core(s["name"]) or (len(x) >= 4 and core(s["name"]) in x))]
        if len(hit) == 1:
            out[hit[0]["id"]] = {"openDate": d, "acquiredDate": "", "note": "売上一覧の「オープン日(譲受日)」", "source": "sales"}
    return out


def from_news(shops):
    news, p = [], 1
    while True:
        try:
            d = get(ES_NEWS.format(p))
        except Exception:
            break
        if not isinstance(d, list) or not d:
            break
        news += d
        p += 1
    brands = sorted({b["name"] for s in shops for b in s.get("brands", [])}, key=len, reverse=True)
    bn = {nrm(b): b for b in brands}
    keys = sorted(list(BRAND_ALIAS) + list(bn), key=len, reverse=True)
    out, unmatched = {}, []
    for n in news:
        t = html.unescape(n["title"]["rendered"])
        if "オープン" not in t:
            continue
        note = "移転新装" if "移転新装" in t else ""
        tt = nrm(re.sub(r"(移転新装)?オープン.*$", "", t))
        brand, rest = None, tt
        for k in keys:
            if tt.startswith(k):
                brand, rest = BRAND_ALIAS.get(k, bn.get(k)), tt[len(k):]
                break
        if not brand:
            unmatched.append((n["date"][:10], t))
            continue
        rest = NAME_ALIAS.get(re.sub(r"店$", "", rest), re.sub(r"店$", "", rest))
        cands = [s for s in shops if any(b["name"] == brand for b in s.get("brands", []))]
        if rest == "" and len(cands) == 1:
            hit = cands
        else:
            hit = [s for s in cands if core(s["name"]) == rest] or [s for s in cands if rest and (rest in core(s["name"]) or core(s["name"]) in rest)]
        if len(hit) == 1:
            sid = hit[0]["id"]
            cur = out.get(sid)
            if not cur or n["date"][:10] < cur["openDate"]:
                out[sid] = {"openDate": n["date"][:10], "acquiredDate": "", "note": (note + "(記事)" if note else ""), "source": "news"}
        else:
            unmatched.append((n["date"][:10], t))
    return out, unmatched


def main():
    shops = load_shops()
    label = {s["id"]: f"{s['brands'][0]['name'] + ' ' if s.get('brands') else ''}{s['name']}" for s in shops}
    lease, lease_file = from_lease(shops)
    sales = from_sales(shops)
    news, news_unmatched = from_news(shops)
    merged = {}
    for src in (news, sales, lease):        # 後勝ち＝lease 最優先
        for sid, v in src.items():
            merged[sid] = v
    payload = {"format": "shop-visit-open-dates-v1", "generated_at": datetime.date.today().isoformat(),
               "sources": {"lease": lease_file, "sales": "ES売上一覧「オープン日(譲受日)」", "news": "eatstyle.jp/news 記事日付"},
               "count": len(merged), "dates": {sid: {**v, "label": label.get(sid, "")} for sid, v in merged.items()}}
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    by = {}
    for v in merged.values():
        by[v["source"]] = by.get(v["source"], 0) + 1
    print(f"wrote {OUT.name}: {len(merged)}/{len(shops)} stores  by source {by}")
    missing = [label[s["id"]] + f"（{s['company']['name']}）" for s in shops if s["id"] not in merged]
    print(f"--- オープン日なし {len(missing)}店（アプリで手入力）")
    for m in missing:
        print(" ", m)
    diff = [(label[sid], lease[sid]["openDate"], sales[sid]["openDate"]) for sid in lease if sid in sales and lease[sid]["openDate"] != sales[sid]["openDate"]]
    print(f"--- lease と売上一覧で日付が異なる {len(diff)}店（lease を採用）")
    for d in diff:
        print(" ", *d)
    print(f"--- ニュース未突合 {len(news_unmatched)}件（閉店・非店舗）: " + " / ".join(t for _, t in news_unmatched))


if __name__ == "__main__":
    main()
